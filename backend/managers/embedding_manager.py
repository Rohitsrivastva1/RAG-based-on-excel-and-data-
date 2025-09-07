"""
Embedding manager for FAISS + LlamaIndex integration.
Handles document creation, indexing, and similarity search.
"""

import os
import pickle
import logging
from typing import List, Dict, Any, Optional, Tuple
from pathlib import Path
import pandas as pd
import numpy as np
from datetime import datetime

# Import our utilities
try:
    from ..settings import settings
    from ..utils.types import Document, QueryResult
    from ..logging_config import get_logger, log_performance
except ImportError:
    from settings import settings
    from utils.types import Document, QueryResult
    from logging_config import get_logger, log_performance

logger = get_logger(__name__)

# Optional imports with fallbacks
try:
    from llama_index.core import Document as LlamaDocument, VectorStoreIndex, StorageContext
    from llama_index.core.base.embeddings.base import BaseEmbedding
    from llama_index.vector_stores.faiss import FaissVectorStore
    from llama_index.embeddings.huggingface import HuggingFaceEmbedding
    from llama_index.core.readers import SimpleDirectoryReader
    LLAMAINDEX_AVAILABLE = True
except ImportError as e:
    LLAMAINDEX_AVAILABLE = False
    logger.warning(f"LlamaIndex not available ({e}). Embedding features will be disabled.")

# Document processing imports
try:
    import fitz  # PyMuPDF
    import docx
    import markdown
    from bs4 import BeautifulSoup
    import requests
    DOCUMENT_PROCESSING_AVAILABLE = True
except ImportError as e:
    DOCUMENT_PROCESSING_AVAILABLE = False
    logger.warning(f"Document processing libraries not available ({e}). Document features will be disabled.")

try:
    import faiss
    FAISS_AVAILABLE = True
except ImportError:
    FAISS_AVAILABLE = False
    logger.warning("FAISS not available. Using simple similarity search.")


class EmbeddingManager:
    """Manages embeddings and vector search using FAISS + LlamaIndex."""
    
    def __init__(self):
        self.embeddings: Optional[BaseEmbedding] = None
        self.vector_stores: Dict[str, FaissVectorStore] = {}
        self.indices: Dict[str, VectorStoreIndex] = {}
        self.documents: Dict[str, List[Document]] = {}
        self.index_dir = Path("./cache/indices")
        self.index_dir.mkdir(parents=True, exist_ok=True)
        
        # Initialize embedding model
        self._initialize_embeddings()
    
    def _initialize_embeddings(self) -> None:
        """Initialize the embedding model."""
        if not LLAMAINDEX_AVAILABLE:
            logger.warning("LlamaIndex not available. Embeddings disabled.")
            return
        
        try:
            if settings.get_embedding_provider() == "huggingface":
                self.embeddings = HuggingFaceEmbedding(
                    model_name=settings.embedding_model,
                    cache_folder=settings.huggingface_cache_dir
                )
                logger.info(f"Initialized HuggingFace embeddings: {settings.embedding_model}")
            else:
                logger.warning("Only HuggingFace embeddings are supported in this version")
                self.embeddings = None
                
        except Exception as e:
            logger.error(f"Failed to initialize embeddings: {e}")
            self.embeddings = None
    
    def create_documents_from_dataframe(
        self, 
        df: pd.DataFrame, 
        session_id: str,
        chunk_size: int = 100
    ) -> List[Document]:
        """
        Create documents from DataFrame for indexing.
        
        Args:
            df: DataFrame to convert
            session_id: Session identifier
            chunk_size: Number of rows per document
            
        Returns:
            List of Document objects
        """
        logger.info(f"Creating documents from DataFrame for session {session_id}", extra={
            "session_id": session_id,
            "rows": len(df),
            "columns": len(df.columns),
            "chunk_size": chunk_size
        })
        
        documents = []
        
        # Create schema document
        schema_doc = Document(
            text=f"Schema: {', '.join(df.columns.tolist())}",
            metadata={
                "type": "schema",
                "session_id": session_id,
                "columns": df.columns.tolist(),
                "dtypes": df.dtypes.to_dict()
            }
        )
        documents.append(schema_doc)
        
        # Create statistics document
        stats_text = f"Dataset statistics: {len(df)} rows, {len(df.columns)} columns. "
        stats_text += f"Columns: {', '.join(df.columns.tolist())}. "
        
        # Add basic statistics for numeric columns
        numeric_cols = df.select_dtypes(include=[np.number]).columns
        if len(numeric_cols) > 0:
            stats_text += f"Numeric columns: {', '.join(numeric_cols)}. "
            for col in numeric_cols[:3]:  # Limit to first 3 numeric columns
                stats_text += f"{col} range: {df[col].min():.2f} to {df[col].max():.2f}. "
        
        stats_doc = Document(
            text=stats_text,
            metadata={
                "type": "statistics",
                "session_id": session_id,
                "row_count": len(df),
                "column_count": len(df.columns)
            }
        )
        documents.append(stats_doc)
        
        # Create row-level documents
        for i in range(0, len(df), chunk_size):
            chunk_df = df.iloc[i:i + chunk_size]
            
            # Create text representation of the chunk
            chunk_text = f"Data rows {i+1}-{min(i+chunk_size, len(df))}: "
            
            # Add column names and sample values
            for col in chunk_df.columns:
                sample_values = chunk_df[col].head(3).tolist()
                chunk_text += f"{col}: {sample_values}. "
            
            # Add aggregated information for numeric columns
            numeric_cols = chunk_df.select_dtypes(include=[np.number]).columns
            if len(numeric_cols) > 0:
                chunk_text += "Aggregates: "
                for col in numeric_cols[:2]:  # Limit to first 2 numeric columns
                    chunk_text += f"{col} sum={chunk_df[col].sum():.2f}, "
                chunk_text = chunk_text.rstrip(", ") + ". "
            
            row_doc = Document(
                text=chunk_text,
                metadata={
                    "type": "data_chunk",
                    "session_id": session_id,
                    "start_row": i,
                    "end_row": min(i + chunk_size, len(df)),
                    "chunk_size": len(chunk_df),
                    "columns": chunk_df.columns.tolist()
                }
            )
            documents.append(row_doc)
        
        # Store documents for this session
        self.documents[session_id] = documents
        
        logger.info(f"Created {len(documents)} documents for session {session_id}")
        return documents
    
    def build_index(self, df: pd.DataFrame, session_id: str) -> bool:
        """
        Build FAISS index for the given DataFrame.
        
        Args:
            df: DataFrame to index
            session_id: Session identifier
            
        Returns:
            True if successful, False otherwise
        """
        if not LLAMAINDEX_AVAILABLE or not self.embeddings:
            logger.warning("Embeddings not available. Index building skipped.")
            return False
        
        if not FAISS_AVAILABLE:
            logger.warning("FAISS not available. Using simple storage.")
            return self._build_simple_index(df, session_id)
        
        try:
            start_time = datetime.utcnow()
            
            # Create documents
            documents = self.create_documents_from_dataframe(df, session_id)
            
            # Convert to LlamaIndex documents
            llama_docs = []
            for doc in documents:
                llama_doc = LlamaDocument(
                    text=doc.text,
                    metadata=doc.metadata
                )
                llama_docs.append(llama_doc)
            
            # Create FAISS vector store
            print(f"🔧 Creating FAISS vector store with dimension: {settings.faiss_dim}")
            faiss_index = faiss.IndexFlatL2(settings.faiss_dim)
            print(f"   FAISS index created: {type(faiss_index)}")
            
            vector_store = FaissVectorStore(faiss_index=faiss_index)
            print(f"   FAISS vector store created: {type(vector_store)}")
            
            # Create storage context
            print(f"🔧 Creating storage context...")
            storage_context = StorageContext.from_defaults(vector_store=vector_store)
            print(f"   Storage context created: {type(storage_context)}")
            
            # Build index
            print(f"🔧 Building VectorStoreIndex with {len(llama_docs)} documents...")
            index = VectorStoreIndex.from_documents(
                llama_docs,
                storage_context=storage_context,
                embed_model=self.embeddings
            )
            print(f"   VectorStoreIndex created: {type(index)}")
            
            # Store index and vector store
            print(f"🔧 Storing index and vector store for session {session_id}...")
            self.indices[session_id] = index
            self.vector_stores[session_id] = vector_store
            print(f"   Index and vector store stored successfully")
            
            # Save index to disk
            print(f"💾 Saving index to disk...")
            self._save_index(session_id, index, vector_store)
            print(f"   Index saved to disk successfully")
            
            duration = (datetime.utcnow() - start_time).total_seconds()
            log_performance("build_index", duration * 1000, 
                          session_id=session_id, 
                          document_count=len(documents))
            
            logger.info(f"Built index for session {session_id}", extra={
                "session_id": session_id,
                "document_count": len(documents),
                "duration_seconds": duration
            })
            
            return True
            
        except Exception as e:
            logger.error(f"Failed to build index for session {session_id}: {e}", exc_info=True)
            return False
    
    def _build_simple_index(self, df: pd.DataFrame, session_id: str) -> bool:
        """Build a simple index without FAISS."""
        try:
            documents = self.create_documents_from_dataframe(df, session_id)
            
            # Store documents in memory
            self.documents[session_id] = documents
            
            logger.info(f"Built simple index for session {session_id}")
            return True
            
        except Exception as e:
            logger.error(f"Failed to build simple index for session {session_id}: {e}")
            return False
    
    def query_index(
        self, 
        session_id: str, 
        query: str, 
        top_k: int = 5
    ) -> List[Dict[str, Any]]:
        """
        Query the index for similar documents.
        
        Args:
            session_id: Session identifier
            query: Query text
            top_k: Number of results to return
            
        Returns:
            List of document results with scores
        """
        if session_id not in self.indices and session_id not in self.documents:
            logger.warning(f"No index found for session {session_id}")
            return []
        
        try:
            if session_id in self.indices:
                # Use FAISS index
                return self._query_faiss_index(session_id, query, top_k)
            else:
                # Use simple search
                return self._query_simple_index(session_id, query, top_k)
                
        except Exception as e:
            logger.error(f"Failed to query index for session {session_id}: {e}")
            return []
    
    def _query_faiss_index(
        self, 
        session_id: str, 
        query: str, 
        top_k: int
    ) -> List[Dict[str, Any]]:
        """Query FAISS index."""
        if session_id not in self.indices:
            return []
        
        try:
            index = self.indices[session_id]
            retriever = index.as_retriever(similarity_top_k=top_k)
            
            # Get similar documents
            nodes = retriever.retrieve(query)
            
            results = []
            for node in nodes:
                results.append({
                    "text": node.text,
                    "metadata": node.metadata,
                    "score": getattr(node, 'score', 0.0)
                })
            
            return results
            
        except Exception as e:
            logger.error(f"Error querying FAISS index: {e}")
            return []
    
    def _query_simple_index(
        self, 
        session_id: str, 
        query: str, 
        top_k: int
    ) -> List[Dict[str, Any]]:
        """Simple text-based search without embeddings."""
        if session_id not in self.documents:
            return []
        
        documents = self.documents[session_id]
        query_lower = query.lower()
        
        # Simple keyword matching
        results = []
        for doc in documents:
            text_lower = doc.text.lower()
            
            # Calculate simple relevance score
            score = 0
            query_words = query_lower.split()
            text_words = text_lower.split()
            
            for word in query_words:
                if word in text_words:
                    score += 1
            
            if score > 0:
                results.append({
                    "text": doc.text,
                    "metadata": doc.metadata,
                    "score": score / len(query_words)  # Normalize score
                })
        
        # Sort by score and return top_k
        results.sort(key=lambda x: x['score'], reverse=True)
        return results[:top_k]
    
    def get_session_info(self, session_id: str) -> Dict[str, Any]:
        """Get information about a session's index."""
        if session_id in self.indices:
            return {
                "enabled": True,
                "document_count": len(self.documents.get(session_id, [])),
                "embedding_model": settings.embedding_model,
                "index_type": "faiss",
                "has_index": True
            }
        elif session_id in self.documents:
            return {
                "enabled": True,
                "document_count": len(self.documents[session_id]),
                "embedding_model": settings.embedding_model,
                "index_type": "simple",
                "has_index": True
            }
        else:
            return {
                "enabled": False,
                "document_count": 0,
                "embedding_model": None,
                "index_type": None,
                "has_index": False
            }
    
    def clear_session(self, session_id: str) -> None:
        """Clear all data for a session."""
        if session_id in self.indices:
            del self.indices[session_id]
        if session_id in self.vector_stores:
            del self.vector_stores[session_id]
        if session_id in self.documents:
            del self.documents[session_id]
        
        # Remove saved index files
        self._remove_index_files(session_id)
        
        logger.info(f"Cleared session {session_id}")
    
    def _save_index(self, session_id: str, index, vector_store) -> None:
        """Save index to disk."""
        try:
            session_dir = self.index_dir / session_id
            session_dir.mkdir(exist_ok=True)
            
            # Save FAISS index
            faiss_path = session_dir / "faiss_index.pkl"
            with open(faiss_path, 'wb') as f:
                # FaissVectorStore stores the index internally, we need to access it differently
                if hasattr(vector_store, 'faiss_index'):
                    pickle.dump(vector_store.faiss_index, f)
                else:
                    # Alternative: save the vector store itself
                    pickle.dump(vector_store, f)
            
            # Save documents
            docs_path = session_dir / "documents.pkl"
            with open(docs_path, 'wb') as f:
                pickle.dump(self.documents.get(session_id, []), f)
            
            logger.info(f"Saved index for session {session_id}")
            
        except Exception as e:
            logger.error(f"Failed to save index for session {session_id}: {e}")
    
    def _load_index(self, session_id: str) -> bool:
        """Load index from disk."""
        try:
            session_dir = self.index_dir / session_id
            
            if not session_dir.exists():
                return False
            
            # Load FAISS index
            faiss_path = session_dir / "faiss_index.pkl"
            if faiss_path.exists():
                with open(faiss_path, 'rb') as f:
                    faiss_index = pickle.load(f)
                
                # Recreate vector store
                vector_store = FaissVectorStore(faiss_index=faiss_index)
                self.vector_stores[session_id] = vector_store
            
            # Load documents
            docs_path = session_dir / "documents.pkl"
            if docs_path.exists():
                with open(docs_path, 'rb') as f:
                    self.documents[session_id] = pickle.load(f)
            
            logger.info(f"Loaded index for session {session_id}")
            return True
            
        except Exception as e:
            logger.error(f"Failed to load index for session {session_id}: {e}")
            return False
    
    def _remove_index_files(self, session_id: str) -> None:
        """Remove index files from disk."""
        try:
            session_dir = self.index_dir / session_id
            if session_dir.exists():
                import shutil
                shutil.rmtree(session_dir)
                logger.info(f"Removed index files for session {session_id}")
        except Exception as e:
            logger.error(f"Failed to remove index files for session {session_id}: {e}")
    
    def process_pdf(self, file_path: str, session_id: str) -> List[Document]:
        """Process PDF file and create documents."""
        if not DOCUMENT_PROCESSING_AVAILABLE:
            raise RuntimeError("Document processing libraries not available")
        
        try:
            logger.info(f"Processing PDF file: {file_path}")
            documents = []
            
            # Open PDF with PyMuPDF
            pdf_doc = fitz.open(file_path)
            
            for page_num in range(pdf_doc.page_count):
                page = pdf_doc[page_num]
                text = page.get_text()
                
                if text.strip():  # Only add non-empty pages
                    doc = Document(
                        text=text,
                        metadata={
                            "type": "pdf_page",
                            "session_id": session_id,
                            "page_number": page_num + 1,
                            "file_path": file_path,
                            "total_pages": pdf_doc.page_count
                        }
                    )
                    documents.append(doc)
            
            pdf_doc.close()
            logger.info(f"Processed PDF: {len(documents)} pages from {file_path}")
            return documents
            
        except Exception as e:
            logger.error(f"Failed to process PDF {file_path}: {e}")
            raise
    
    def process_docx(self, file_path: str, session_id: str) -> List[Document]:
        """Process DOCX file and create documents."""
        if not DOCUMENT_PROCESSING_AVAILABLE:
            raise RuntimeError("Document processing libraries not available")
        
        try:
            logger.info(f"Processing DOCX file: {file_path}")
            documents = []
            
            # Open DOCX with python-docx
            doc = docx.Document(file_path)
            
            # Extract text from paragraphs
            full_text = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    full_text.append(paragraph.text.strip())
            
            # Split into chunks (e.g., every 500 words)
            chunk_size = 500
            text = " ".join(full_text)
            words = text.split()
            
            for i in range(0, len(words), chunk_size):
                chunk_words = words[i:i + chunk_size]
                chunk_text = " ".join(chunk_words)
                
                doc_obj = Document(
                    text=chunk_text,
                    metadata={
                        "type": "docx_chunk",
                        "session_id": session_id,
                        "chunk_index": i // chunk_size + 1,
                        "file_path": file_path,
                        "word_count": len(chunk_words)
                    }
                )
                documents.append(doc_obj)
            
            logger.info(f"Processed DOCX: {len(documents)} chunks from {file_path}")
            return documents
            
        except Exception as e:
            logger.error(f"Failed to process DOCX {file_path}: {e}")
            raise
    
    def process_markdown(self, file_path: str, session_id: str) -> List[Document]:
        """Process Markdown file and create documents."""
        if not DOCUMENT_PROCESSING_AVAILABLE:
            raise RuntimeError("Document processing libraries not available")
        
        try:
            logger.info(f"Processing Markdown file: {file_path}")
            documents = []
            
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
            
            # Convert markdown to HTML then extract text
            html = markdown.markdown(content)
            soup = BeautifulSoup(html, 'html.parser')
            text = soup.get_text()
            
            # Split into chunks (e.g., every 300 words)
            chunk_size = 300
            words = text.split()
            
            for i in range(0, len(words), chunk_size):
                chunk_words = words[i:i + chunk_size]
                chunk_text = " ".join(chunk_words)
                
                doc = Document(
                    text=chunk_text,
                    metadata={
                        "type": "markdown_chunk",
                        "session_id": session_id,
                        "chunk_index": i // chunk_size + 1,
                        "file_path": file_path,
                        "word_count": len(chunk_words)
                    }
                )
                documents.append(doc)
            
            logger.info(f"Processed Markdown: {len(documents)} chunks from {file_path}")
            return documents
            
        except Exception as e:
            logger.error(f"Failed to process Markdown {file_path}: {e}")
            raise
    
    def process_website(self, url: str, session_id: str) -> List[Document]:
        """Process website URL and create documents."""
        if not DOCUMENT_PROCESSING_AVAILABLE:
            raise RuntimeError("Document processing libraries not available")
        
        try:
            logger.info(f"Processing website: {url}")
            documents = []
            
            # Fetch webpage content
            response = requests.get(url, timeout=30)
            response.raise_for_status()
            
            # Parse HTML content
            soup = BeautifulSoup(response.content, 'html.parser')
            
            # Remove script and style elements
            for script in soup(["script", "style"]):
                script.decompose()
            
            # Extract text
            text = soup.get_text()
            
            # Clean up text
            lines = (line.strip() for line in text.splitlines())
            chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
            text = ' '.join(chunk for chunk in chunks if chunk)
            
            # Split into chunks (e.g., every 400 words)
            chunk_size = 400
            words = text.split()
            
            for i in range(0, len(words), chunk_size):
                chunk_words = words[i:i + chunk_size]
                chunk_text = " ".join(chunk_words)
                
                doc = Document(
                    text=chunk_text,
                    metadata={
                        "type": "website_chunk",
                        "session_id": session_id,
                        "chunk_index": i // chunk_size + 1,
                        "url": url,
                        "word_count": len(chunk_words)
                    }
                )
                documents.append(doc)
            
            logger.info(f"Processed website: {len(documents)} chunks from {url}")
            return documents
            
        except Exception as e:
            logger.error(f"Failed to process website {url}: {e}")
            raise
    
    def add_documents_to_index(self, documents: List[Document], session_id: str) -> bool:
        """Add documents to the existing index for a session."""
        if not LLAMAINDEX_AVAILABLE or not self.embeddings:
            logger.warning("Embeddings not available. Document indexing skipped.")
            return False
        
        try:
            # Convert to LlamaIndex documents
            llama_docs = []
            for doc in documents:
                llama_doc = LlamaDocument(
                    text=doc.text,
                    metadata=doc.metadata
                )
                llama_docs.append(llama_doc)
            
            if session_id in self.indices:
                # Add to existing index
                index = self.indices[session_id]
                for llama_doc in llama_docs:
                    index.insert(llama_doc)
                logger.info(f"Added {len(documents)} documents to existing index for session {session_id}")
            else:
                # Create new index
                if not FAISS_AVAILABLE:
                    logger.warning("FAISS not available. Using simple storage.")
                    return self._add_documents_simple(documents, session_id)
                
                # Create FAISS vector store
                faiss_index = faiss.IndexFlatL2(settings.faiss_dim)
                vector_store = FaissVectorStore(faiss_index=faiss_index)
                storage_context = StorageContext.from_defaults(vector_store=vector_store)
                
                # Build index
                index = VectorStoreIndex.from_documents(
                    llama_docs,
                    storage_context=storage_context,
                    embed_model=self.embeddings
                )
                
                # Store index and vector store
                self.indices[session_id] = index
                self.vector_stores[session_id] = vector_store
                
                logger.info(f"Created new index with {len(documents)} documents for session {session_id}")
            
            # Save index to disk
            self._save_index(session_id, self.indices[session_id], self.vector_stores[session_id])
            
            return True
            
        except Exception as e:
            logger.error(f"Failed to add documents to index for session {session_id}: {e}")
            return False
    
    def _add_documents_simple(self, documents: List[Document], session_id: str) -> bool:
        """Add documents using simple storage (no FAISS)."""
        try:
            if session_id not in self.documents:
                self.documents[session_id] = []
            
            self.documents[session_id].extend(documents)
            logger.info(f"Added {len(documents)} documents to simple storage for session {session_id}")
            return True
            
        except Exception as e:
            logger.error(f"Failed to add documents to simple storage for session {session_id}: {e}")
            return False


# Global embedding manager instance
embedding_manager: Optional[EmbeddingManager] = None


def initialize_embedding_manager() -> EmbeddingManager:
    """Initialize the global embedding manager."""
    global embedding_manager
    embedding_manager = EmbeddingManager()
    return embedding_manager


def get_embedding_manager() -> EmbeddingManager:
    """Get the global embedding manager instance."""
    if embedding_manager is None:
        raise RuntimeError("Embedding manager not initialized. Call initialize_embedding_manager() first.")
    return embedding_manager


# Example usage and testing
if __name__ == "__main__":
    import pandas as pd
    
    # Test embedding manager
    manager = EmbeddingManager()
    
    # Create test data
    df = pd.DataFrame({
        'Category': ['Education', 'Entertainment', 'Finance'],
        'Revenue': [1000, 1500, 1200],
        'Users': [100, 150, 120]
    })
    
    # Test document creation
    session_id = "test-session"
    documents = manager.create_documents_from_dataframe(df, session_id)
    print(f"Created {len(documents)} documents")
    
    # Test index building
    success = manager.build_index(df, session_id)
    print(f"Index building successful: {success}")
    
    # Test querying
    results = manager.query_index(session_id, "revenue by category", top_k=3)
    print(f"Query results: {len(results)} documents")
    
    # Test session info
    info = manager.get_session_info(session_id)
    print(f"Session info: {info}")
    
    # Cleanup
    manager.clear_session(session_id)
    print("Session cleared")
